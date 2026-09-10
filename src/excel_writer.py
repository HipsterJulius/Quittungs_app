"""
Schreibt extrahierte Bon-Daten in die Haushaltsbuch-Excel-Datei.

Jede Zeile = ein Produkt. Spalten: Einkaufs-ID, Produkt, Betrag, Kategorie,
Datum, Gesamtbetrag, Supermarktkette.

Die Einkaufs-ID ist eine automatisch hochzählende Nummer: alle Produkte
desselben Bons bekommen dieselbe ID, der nächste Bon die nächsthöhere -
damit lassen sich Einkäufe in Excel z.B. per Pivot-Tabelle sauber trennen.

Falls die Excel-Datei schon existiert, werden neue Zeilen angehängt statt
die Datei zu überschreiben - so sammelt sich das Haushaltsbuch über
mehrere Bons hinweg.

Voraussetzungen:
    pip install openpyxl
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook

COLUMNS = ["Einkaufs-ID", "Produkt", "Betrag", "Kategorie", "Datum", "Gesamtbetrag", "Supermarktkette"]


def _ensure_id_column(ws) -> None:
    """
    Migration für Dateien aus einer älteren Skript-Version ohne
    Einkaufs-ID-Spalte: fügt sie nachträglich als erste Spalte ein.
    Bereits vorhandene Zeilen (aus Testläufen vor dieser Änderung)
    bekommen alle die ID 1, da sie ohnehin nicht unterscheidbar waren.
    """
    header = [cell.value for cell in ws[1]]
    if header and header[0] == "Einkaufs-ID":
        return  # Datei ist schon aktuell

    ws.insert_cols(1)
    ws.cell(row=1, column=1, value="Einkaufs-ID")
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=1, value=1)


def _next_purchase_id(ws) -> int:
    """Höchste vorhandene Einkaufs-ID + 1. Bei leerer Tabelle: 1."""
    max_id = 0
    for (value,) in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if isinstance(value, (int, float)):
            max_id = max(max_id, int(value))
    return max_id + 1


def append_receipt(data: dict, excel_path: str) -> int:
    """
    data: das dict aus llm_extractor.extract_structured_data()
          (Schlüssel: haendler, datum, gesamtbetrag, produkte)
    excel_path: Pfad zur Haushaltsbuch-Excel-Datei (wird angelegt, falls
                sie noch nicht existiert)

    Gibt die vergebene Einkaufs-ID zurück.
    """
    path = Path(excel_path)
    path.parent.mkdir(parents=True, exist_ok=True)  # z.B. output/-Ordner anlegen, falls nötig

    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
        _ensure_id_column(ws)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Haushaltsbuch"
        ws.append(COLUMNS)

    purchase_id = _next_purchase_id(ws)

    haendler = data["haendler"]
    datum = data["datum"]
    gesamtbetrag = data["gesamtbetrag"]

    for produkt in data["produkte"]:
        ws.append([
            purchase_id,
            produkt["name"],
            produkt["preis"],
            produkt["kategorie"],
            datum,
            gesamtbetrag,
            haendler,
        ])

    wb.save(path)
    return purchase_id


if __name__ == "__main__":
    import json
    import sys

    if len(sys.argv) < 3:
        print("Nutzung: python excel_writer.py <json_datei> <ausgabe.xlsx>")
        sys.exit(1)

    with open(sys.argv[1], "r", encoding="utf-8") as f:
        data = json.load(f)

    purchase_id = append_receipt(data, sys.argv[2])
    print(f"Geschrieben nach: {sys.argv[2]} (Einkaufs-ID: {purchase_id})")
